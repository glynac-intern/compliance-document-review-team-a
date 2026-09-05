"""
Extracts plain text from uploaded documents, per Document.type.

Used by the analysis pipeline before PII masking -- masking must run on
extracted text, not the raw file, since a PDF/DOCX's binary structure
isn't something the masker's regexes can operate on directly.
"""

from pathlib import Path

from pypdf import PdfReader
from docx import Document as DocxDocument
from openpyxl import load_workbook


def extract_pdf(file_path: str) -> str:
    reader = PdfReader(file_path)
    return "\n\n".join(page.extract_text() or "" for page in reader.pages)


def extract_docx(file_path: str) -> str:
    doc = DocxDocument(file_path)
    return "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())


def extract_xlsx(file_path: str) -> str:
    wb = load_workbook(file_path, data_only=True)
    parts = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            row_text = " | ".join(str(cell) for cell in row if cell is not None)
            if row_text.strip():
                parts.append(row_text)
    return "\n".join(parts)


EXTRACTORS = {
    "pdf": extract_pdf,
    "docx": extract_docx,
    "xlsx": extract_xlsx,
}


def extract_text(file_path: str, doc_type: str) -> str:
    """doc_type is the string value of models.DocumentType (e.g. 'pdf')."""
    extractor = EXTRACTORS.get(doc_type)
    if extractor is None:
        raise ValueError(f"No extractor for document type: {doc_type}")
    text = extractor(file_path)
    if not text.strip():
        raise ValueError(f"Extraction produced no text from {file_path} (type={doc_type})")
    return text
