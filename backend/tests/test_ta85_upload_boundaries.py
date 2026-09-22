"""
Tests for TA-85: upload validation at its boundaries. Upload validation
is the app's outermost input boundary, and nothing exercised it fully --
test_file_signature.py (TA-23) and test_upload_size_cap.py (TA-24)
already cover mislabelled-content-type rejection and the streaming size
cap at the function level, but nothing covered: every accepted format
actually being accepted end to end, a genuinely unrecognized format
being rejected (a different code path than "mislabelled" -- here
_detect_file_type finds no match at all, not a mismatch), the size cap
enforced through the real endpoint (not just the helper function), or
what happens to a file that looks like an allowed type but is corrupt
once something actually tries to read it.
"""

import pytest

pytestmark = pytest.mark.integration

import io

from docx import Document as DocxDocument
from openpyxl import Workbook
from pypdf import PdfWriter

from documents.router import MAX_FILE_SIZE
from models import AIAnalysis, AnalysisStatus, DocumentChunk, Flag, PIIMapping

CORRUPT_BUT_SIGNATURE_VALID_PDF = (
    "corrupt.pdf",
    b"%PDF-1.4 minimal fake content",  # the fixture used throughout this
    # test suite as "a PDF" -- it satisfies _detect_file_type's magic-byte
    # check, but is not a structurally valid PDF: pypdf raises reading it.
    "application/pdf",
)


def _real_pdf_bytes() -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def _real_docx_bytes() -> bytes:
    doc = DocxDocument()
    doc.add_paragraph("This is a real, valid, parseable DOCX document.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _real_xlsx_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value="This is a real, valid, parseable XLSX document.")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_pdf_is_accepted(client, advisor_token):
    resp = client.post(
        "/documents",
        headers={"Authorization": f"Bearer {advisor_token}"},
        files={"file": ("real.pdf", _real_pdf_bytes(), "application/pdf")},
    )
    assert resp.status_code == 201, resp.text
    assert resp.json()["type"] == "pdf"


def test_docx_is_accepted(client, advisor_token):
    resp = client.post(
        "/documents",
        headers={"Authorization": f"Bearer {advisor_token}"},
        files={
            "file": (
                "real.docx",
                _real_docx_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert resp.status_code == 201, resp.text
    assert resp.json()["type"] == "docx"


def test_xlsx_is_accepted(client, advisor_token):
    resp = client.post(
        "/documents",
        headers={"Authorization": f"Bearer {advisor_token}"},
        files={
            "file": (
                "real.xlsx",
                _real_xlsx_bytes(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert resp.status_code == 201, resp.text
    assert resp.json()["type"] == "xlsx"


def test_unrecognized_format_is_rejected(client, advisor_token):
    """A different failure path than "mislabelled" (TA-23's existing
    coverage): here the file's signature doesn't match ANY supported
    type at all, not just a mismatch against the declared one."""
    resp = client.post(
        "/documents",
        headers={"Authorization": f"Bearer {advisor_token}"},
        files={"file": ("notes.txt", b"Just plain text, no known signature at all.", "text/plain")},
    )
    assert resp.status_code == 400
    assert "unsupported or unrecognized" in resp.json()["detail"].lower()


@pytest.mark.large_upload
def test_file_exactly_at_the_size_cap_is_accepted_through_the_real_endpoint(client, advisor_token):
    content = b"%PDF-" + b"0" * (MAX_FILE_SIZE - 5)
    assert len(content) == MAX_FILE_SIZE

    resp = client.post(
        "/documents",
        headers={"Authorization": f"Bearer {advisor_token}"},
        files={"file": ("at-cap.pdf", content, "application/pdf")},
    )
    assert resp.status_code == 201, resp.text


@pytest.mark.large_upload
def test_file_one_byte_over_the_size_cap_is_rejected_through_the_real_endpoint(client, advisor_token):
    content = b"%PDF-" + b"0" * (MAX_FILE_SIZE - 5 + 1)
    assert len(content) == MAX_FILE_SIZE + 1

    resp = client.post(
        "/documents",
        headers={"Authorization": f"Bearer {advisor_token}"},
        files={"file": ("over-cap.pdf", content, "application/pdf")},
    )
    assert resp.status_code == 400
    assert "10mb" in resp.json()["detail"].lower()


def test_corrupt_file_of_an_allowed_type_fails_analysis_cleanly_not_a_crash(
    client, db_session, advisor_token
):
    """Upload validation only checks the file's magic-byte signature, not
    that it's a fully well-formed document -- so a corrupt file of an
    allowed type is accepted at upload time, same as every other test in
    this suite that uses this exact fixture. What matters is that
    actually trying to analyze it fails cleanly (503, status=failed) via
    _execute_analysis's existing exception handling, instead of a raw
    500 or a hang -- and leaves no partial Flag/DocumentChunk/PIIMapping
    rows behind. Needs no LLM mocking at all: extraction fails before
    any vendor call would even happen."""
    submit = client.post(
        "/documents",
        headers={"Authorization": f"Bearer {advisor_token}"},
        files={"file": CORRUPT_BUT_SIGNATURE_VALID_PDF},
    )
    assert submit.status_code == 201, submit.text
    doc_id = submit.json()["id"]

    resp = client.get(
        f"/documents/{doc_id}/analysis",
        headers={"Authorization": f"Bearer {advisor_token}"},
    )
    assert resp.status_code == 503

    analysis = db_session.query(AIAnalysis).filter(AIAnalysis.document_id == doc_id).first()
    assert analysis.status == AnalysisStatus.failed
    assert "PdfStreamError" in analysis.error_message or "stream" in analysis.error_message.lower()

    assert db_session.query(Flag).filter(Flag.analysis_id == analysis.id).count() == 0
    assert db_session.query(DocumentChunk).filter(DocumentChunk.document_id == doc_id).count() == 0
    assert db_session.query(PIIMapping).filter(PIIMapping.document_id == doc_id).count() == 0
